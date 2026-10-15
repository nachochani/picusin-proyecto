from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from app.security import get_current_admin
from app.database import get_db
from app.models.product import Product, EstadoProducto
from app.models.order import Order
from app.models.reservation import Reservation
from app.models.discount_code import Discount_code
import openpyxl
import io
import re
import html
from html.parser import HTMLParser
from urllib.request import Request, urlopen
from urllib.parse import urljoin, urlparse


router = APIRouter(
    dependencies=[Depends(get_current_admin)]
)


class BuscadorImagen(HTMLParser):
    def __init__(self):
        super().__init__()
        self.imagen = None

    def handle_starttag(self, tag, attrs):
        if tag.lower() != "meta":
            return

        atributos = dict(attrs)
        propiedad = (
            atributos.get("property")
            or atributos.get("name")
            or ""
        ).lower()

        if propiedad in ("og:image", "twitter:image"):
            contenido = atributos.get("content")
            if contenido and not self.imagen:
                self.imagen = contenido


def obtener_url_imagen(url):
    if not url:
        return None

    url = str(url).strip()

    # Si el enlace ya apunta a un archivo de imagen,
    # lo guardamos directamente.
    extensiones = (".jpg", ".jpeg", ".png", ".webp", ".gif")

    ruta = urlparse(url).path.lower()
    if ruta.endswith(extensiones):
        return url

    # Si es una página de producto, buscamos sus metadatos.
    try:
        solicitud = Request(
            url,
            headers={"User-Agent": "Mozilla/5.0"}
        )

        with urlopen(solicitud, timeout=5) as respuesta:
            contenido = respuesta.read(2_000_000).decode(
                "utf-8", errors="ignore"
            )

        buscador = BuscadorImagen()
        buscador.feed(contenido)

        if buscador.imagen:
            imagen = html.unescape(buscador.imagen)
            return urljoin(url, imagen)

    except Exception:
        pass

    return None


@router.get("/productos")
def get_productos(db: Session = Depends(get_db)):
    productos = db.query(Product).filter(
        Product.activo == True,
    ).all()
    return productos

@router.post("/productos")
def crear_producto(
    nombre:str,
    precio: float,
    imagen:str = None,
    es_novedad: bool = False,
    es_preventa: bool = False,
    db: Session = Depends(get_db)
    ):
    nuevo_producto = Product(
        nombre=nombre.strip(),
        precio=precio,
        estado = EstadoProducto.disponible,
        es_preventa=es_preventa,
        es_novedad=es_novedad,
        imagen=imagen,
        activo=True
    )  
    db.add(nuevo_producto)
    db.commit()
    db.refresh(nuevo_producto)

    return nuevo_producto

@router.put("/productos/{id}")
def update_product(
    id: int,
    nombre: str = None,
    precio: float = None,
    estado: str = None,
    imagen: str = None,
    db: Session = Depends(get_db)):
    producto = db.query(Product).filter(Product.id == id).first()

    if not producto:
        raise HTTPException(status_code=404, detail = "Producto no encontrado")

    if nombre: producto.nombre = nombre
    if precio: producto.precio = precio
    if estado: producto.estado = estado
    if imagen: producto.imagen = imagen


    db.commit()
    db.refresh(producto)

    return producto

@router.delete("/productos/{id}")
def delete_product(id: int, db: Session = Depends(get_db)):
    producto = db.query(Product).filter(Product.id == id).first()

    if not producto:
        raise HTTPException(status_code=404, detail="Producto no encontrado")

    producto.activo = False

    db.commit()
    db.refresh(producto)

    return producto


@router.post("/cargar-excel")
async def cargar_excel(
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    contenido = await archivo.read()
    wb = openpyxl.load_workbook(
        io.BytesIO(contenido),
        data_only=True
    )

    nombre_catalogo = "📦 CATALOGO DE PRODUCTOS"
    nombre_pedidos = "ARTICULOS PARA TRAER A PEDIDO"

    pestañas = {
        nombre_catalogo: False,
        nombre_pedidos: True,
    }

    productos_cargados = 0
    productos_omitidos = 0
    imagenes_pendientes = 0

    def convertir_precio(valor):
        if isinstance(valor, (int, float)):
            return float(valor)

        precio_str = str(valor).replace("$", "").replace(
            " ", ""
        ).strip()

        if not precio_str:
            return None

        if "," in precio_str and "." in precio_str:
            if precio_str.rfind(",") > precio_str.rfind("."):
                precio_str = precio_str.replace(".", "")
                precio_str = precio_str.replace(",", ".")
            else:
                precio_str = precio_str.replace(",", "")
        elif "," in precio_str:
            precio_str = precio_str.replace(",", ".")
        elif "." in precio_str:
            partes = precio_str.split(".")
            if len(partes[-1]) == 3:
                precio_str = precio_str.replace(".", "")

        try:
            return float(precio_str)
        except (ValueError, TypeError):
            return None

    try:
        for nombre_hoja, es_novedad in pestañas.items():
            if nombre_hoja not in wb.sheetnames:
                raise HTTPException(
                    status_code=400,
                    detail=f"No se encontró la pestaña: {nombre_hoja}"
                )

            hoja = wb[nombre_hoja]
            es_pestana_pedidos = (
                nombre_hoja == nombre_pedidos
            )

            # El catálogo general tiene una fila adicional
            # de categoría antes de los artículos.
            fila_inicial = 4 if not es_pestana_pedidos else 3

            for fila in hoja.iter_rows(
                min_row=fila_inicial,
                values_only=False
            ):
                celda_nombre = fila[0]
                celda_precio = fila[1] if len(fila) > 1 else None
                celda_estado = fila[2] if len(fila) > 2 else None
                celda_imagen = fila[3] if len(fila) > 3 else None

                nombre = celda_nombre.value
                precio_original = (
                    celda_precio.value if celda_precio else None
                )

                # Ignorar filas vacías o sin un precio válido.
                if not nombre or precio_original is None:
                    continue

                precio = convertir_precio(precio_original)

                if precio is None or precio <= 0:
                    productos_omitidos += 1
                    continue

                nombre = str(nombre).strip()

                estado_excel = (
                    str(celda_estado.value).upper()
                    if celda_estado and celda_estado.value
                    else ""
                )

                marcado_stock = "STOCK" in estado_excel
                marcado_pedido = "A PEDIDO" in estado_excel

                es_preventa = (
                    es_pestana_pedidos or marcado_pedido
                )

                # Para los artículos a pedido no asignamos
                # unidades físicas de stock.
                tiene_stock = marcado_stock and not es_preventa

                stock = 1 if tiene_stock else 0
                estado = EstadoProducto.disponible

                # Evitar duplicados al volver a subir el mismo Excel.
                existente = db.query(Product).filter(
                    Product.nombre == nombre,
                    Product.precio == precio,
                    Product.es_novedad == es_novedad,
                    Product.es_preventa == es_preventa,
                ).first()

                if existente:
                    productos_omitidos += 1
                    continue

                imagen_url = None

                if celda_imagen:
                    if celda_imagen.hyperlink:
                        enlace = celda_imagen.hyperlink.target
                    else:
                        valor_imagen = celda_imagen.value
                        enlace = (
                            str(valor_imagen).strip()
                            if valor_imagen
                            and str(valor_imagen).strip().startswith(
                                ("http://", "https://")
                            )
                            else None
                        )

                    if enlace:
                        imagen_url = obtener_url_imagen(enlace)

                if not imagen_url:
                    imagenes_pendientes += 1

                nuevo_producto = Product(
                    nombre=nombre,
                    descripcion=nombre,
                    precio=precio,
                    stock=stock,
                    estado=estado,
                    es_preventa=es_preventa,
                    es_novedad=es_novedad,
                    imagen=imagen_url,
                    activo=True,
                )

                db.add(nuevo_producto)
                productos_cargados += 1

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "mensaje": "Carga completada",
        "productos_cargados": productos_cargados,
        "productos_omitidos": productos_omitidos,
        "imagenes_pendientes": imagenes_pendientes,
    }


@router.get("/reservas")
def get_reservas(db: Session = Depends(get_db)):
    reservas = db.query(Reservation).all()
    return reservas

@router.put("/reservas/{id}/estado")
def update_estado_reserva(id: int, estado: str, db: Session = Depends(get_db)):
    reserva = db.query(Reservation).filter(Reservation.id == id).first()

    if not reserva:
        raise HTTPException(status_code=404, detail = "Reserva no encontrada")
    estados_validos = ["pendiente", "confirmada", "cancelada"]

    if estado not in estados_validos:
        raise HTTPException(status_code=400, detail = "Estado no valido")
    reserva.estado = estado

    db.commit()
    db.refresh(reserva)

    return reserva


@router.get("/descuentos")
def get_descuentos(db: Session = Depends(get_db)):
    descuentos = db.query(Discount_code).order_by(
        Discount_code.id.desc()
    ).all()

    return descuentos


@router.post("/descuentos")
def crear_descuento(
    codigo: str,
    porcentaje: float,
    db: Session = Depends(get_db)
):
    codigo = codigo.strip().upper()

    if not codigo:
        raise HTTPException(
            status_code=400,
            detail="El código no puede estar vacío"
        )

    if porcentaje <= 0 or porcentaje > 100:
        raise HTTPException(
            status_code=400,
            detail="El porcentaje debe ser mayor que 0 y menor o igual a 100"
        )

    existente = db.query(Discount_code).filter(
        Discount_code.codigo == codigo
    ).first()

    if existente:
        raise HTTPException(
            status_code=400,
            detail="Ya existe un código con ese nombre"
        )

    nuevo_descuento = Discount_code(
        codigo=codigo,
        porcentaje=porcentaje,
        activo=True
    )

    db.add(nuevo_descuento)
    db.commit()
    db.refresh(nuevo_descuento)

    return nuevo_descuento


@router.put("/descuentos/{id}")
def actualizar_descuento(
    id: int,
    codigo: str,
    porcentaje: float,
    db: Session = Depends(get_db)
):
    descuento = db.query(Discount_code).filter(
        Discount_code.id == id
    ).first()

    if not descuento:
        raise HTTPException(
            status_code=404,
            detail="Código de descuento no encontrado"
        )

    codigo = codigo.strip().upper()

    if not codigo:
        raise HTTPException(
            status_code=400,
            detail="El código no puede estar vacío"
        )

    if porcentaje <= 0 or porcentaje > 100:
        raise HTTPException(
            status_code=400,
            detail="El porcentaje debe ser mayor que 0 y menor o igual a 100"
        )

    existente = db.query(Discount_code).filter(
        Discount_code.codigo == codigo,
        Discount_code.id != id
    ).first()

    if existente:
        raise HTTPException(
            status_code=400,
            detail="Ya existe otro código con ese nombre"
        )

    descuento.codigo = codigo
    descuento.porcentaje = porcentaje

    db.commit()
    db.refresh(descuento)

    return descuento


@router.put("/descuentos/{id}/estado")
def actualizar_estado_descuento(
    id: int,
    activo: bool,
    db: Session = Depends(get_db)
):
    descuento = db.query(Discount_code).filter(
        Discount_code.id == id
    ).first()

    if not descuento:
        raise HTTPException(
            status_code=404,
            detail="Código de descuento no encontrado"
        )

    descuento.activo = activo

    db.commit()
    db.refresh(descuento)

    return descuento


@router.delete("/descuentos/{id}")
def eliminar_descuento(
    id: int,
    db: Session = Depends(get_db)
):
    descuento = db.query(Discount_code).filter(
        Discount_code.id == id
    ).first()

    if not descuento:
        raise HTTPException(
            status_code=404,
            detail="Código de descuento no encontrado"
        )

    db.delete(descuento)
    db.commit()

    return {
        "mensaje": "Código de descuento eliminado correctamente"
    }
